"""Local, single-user document store for the native desktop application."""
from __future__ import annotations

import csv
import getpass
import hashlib
import json
import os
import re
import shutil
import sqlite3
import tempfile
import zipfile
from collections.abc import Iterator
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from docket import catalog
from docket.export.tabular import (
    ITEM_COLUMNS,
    RESULT_COLUMNS,
    line_item_rows,
    result_row,
)
from docket.result import DocumentResult
from docket.validate import validate

SUPPORTED = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
DOCUMENT_COLUMNS = RESULT_COLUMNS + (
    "filename", "stored_path", "created_at", "updated_at", "approval_status", "approved_at",
)
DETAIL_COLUMNS = ("document_id", "path", "value_type", "part", "value")
HISTORY_COLUMNS = ("document_id", "revision_id", "at", "action", "part", "detail_json")
MAX_CELL_CHARS = 16_000


def default_data_dir() -> Path:
    import os
    import sys

    if location := os.environ.get("DOCKET_DESKTOP_DATA_DIR"):
        return Path(location).expanduser()
    if sys.platform == "darwin":
        return Path.home() / "Library" / "Application Support" / "Docket Desktop"
    if sys.platform == "win32":
        return Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "Docket Desktop"
    return Path(os.environ.get("XDG_DATA_HOME", str(Path.home() / ".local" / "share"))) / "docket-desktop"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _csv_safe(value: Any) -> Any:
    """Keep untrusted document text from becoming a spreadsheet formula."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_safe(value: Any) -> Any:
    value = _csv_safe(value)
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    return value


def _document_row(row: dict[str, Any], result: DocumentResult) -> dict[str, Any]:
    return result_row(result) | {
        "filename": row["filename"],
        "stored_path": row["stored_path"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
        "approval_status": "approved",
        "approved_at": row["approved_at"],
    }


def _detail_rows(document_id: str, value: Any, path: str = "") -> Iterator[dict[str, Any]]:
    """Flatten nested data without losing keys or overfilling spreadsheet cells."""
    if isinstance(value, dict) and value:
        for key, child in value.items():
            yield from _detail_rows(document_id, child,
                                    path + "/" + str(key).replace("~", "~0").replace("/", "~1"))
        return
    if isinstance(value, list) and value:
        for index, child in enumerate(value):
            yield from _detail_rows(document_id, child, path + f"/{index}")
        return
    if isinstance(value, dict):
        value_type, rendered = "object", "{}"
    elif isinstance(value, list):
        value_type, rendered = "array", "[]"
    elif value is None:
        value_type, rendered = "null", ""
    elif isinstance(value, bool):
        value_type, rendered = "boolean", "true" if value else "false"
    elif isinstance(value, int):
        value_type, rendered = "integer", str(value)
    elif isinstance(value, float):
        value_type, rendered = "number", repr(value)
    else:
        value_type, rendered = "string", str(value)
    for part, start in enumerate(range(0, max(1, len(rendered)), MAX_CELL_CHARS), start=1):
        yield {"document_id": document_id, "path": path or "/", "value_type": value_type,
               "part": part, "value": rendered[start:start + MAX_CELL_CHARS]}


def _changes(before: Any, after: Any, path: str = "") -> list[dict]:
    if isinstance(before, dict) and isinstance(after, dict):
        return [change for key in sorted(before.keys() | after.keys())
                for change in _changes(before.get(key), after.get(key),
                                       f"{path}.{key}" if path else key)]
    if isinstance(before, list) and isinstance(after, list):
        return [change for index in range(max(len(before), len(after)))
                for change in _changes(before[index] if index < len(before) else None,
                                       after[index] if index < len(after) else None,
                                       f"{path}[{index}]")]
    return [{"field": path, "before": before, "after": after}] if before != after else []


class DesktopStore:
    def __init__(self, root: Path):
        self.root = Path(root)
        self.files = self.root / "files"
        self.root.mkdir(parents=True, exist_ok=True)
        self.files.mkdir(exist_ok=True)
        if os.name != "nt":
            self.root.chmod(0o700)
            self.files.chmod(0o700)
        self.db = sqlite3.connect(self.root / "documents.sqlite3")
        if os.name != "nt":
            (self.root / "documents.sqlite3").chmod(0o600)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys=ON")
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.executescript("""
            CREATE TABLE IF NOT EXISTS documents (
                id TEXT PRIMARY KEY, filename TEXT NOT NULL, stored_path TEXT NOT NULL,
                kind TEXT, state TEXT NOT NULL, result_json TEXT, edited_json TEXT,
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL, approved_at TEXT
            );
            CREATE TABLE IF NOT EXISTS revisions (
                id INTEGER PRIMARY KEY AUTOINCREMENT, document_id TEXT NOT NULL,
                at TEXT NOT NULL, action TEXT NOT NULL, detail_json TEXT,
                FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
            );
        """)
        self.db.commit()

    def close(self) -> None:
        self.db.close()

    def _event(self, document_id: str, action: str, detail: dict | None = None) -> None:
        self.db.execute(
            "INSERT INTO revisions(document_id, at, action, detail_json) VALUES(?,?,?,?)",
            (document_id, _now(), action, json.dumps(detail or {}, ensure_ascii=False)),
        )

    def add_file(self, path: Path) -> tuple[str, bool]:
        path = Path(path)
        if path.suffix.lower() not in SUPPORTED or not path.is_file():
            raise ValueError("Choose a PDF or image file")
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            while block := handle.read(1 << 20):
                digest.update(block)
        document_id = digest.hexdigest()
        if self.get(document_id):
            return document_id, False
        stored = self.files / f"{document_id}{path.suffix.lower()}"
        shutil.copyfile(path, stored)
        if os.name != "nt":
            stored.chmod(0o600)
        now = _now()
        with self.db:
            self.db.execute(
                "INSERT INTO documents(id,filename,stored_path,state,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                (document_id, path.name, str(stored), "new", now, now),
            )
            self._event(document_id, "imported", {"filename": path.name})
        return document_id, True

    def get(self, document_id: str) -> dict | None:
        row = self.db.execute("SELECT * FROM documents WHERE id=?", (document_id,)).fetchone()
        return dict(row) if row else None

    def list_documents(self, query: str = "") -> list[dict]:
        pattern = f"%{query.strip()}%"
        return [dict(row) for row in self.db.execute(
            "SELECT * FROM documents WHERE filename LIKE ? OR kind LIKE ? OR state LIKE ? ORDER BY created_at DESC, id DESC",
            (pattern, pattern, pattern),
        )]

    def history(self, document_id: str) -> list[dict]:
        return [dict(row) for row in self.db.execute(
            "SELECT id,at,action,detail_json FROM revisions WHERE document_id=? ORDER BY id",
            (document_id,),
        )]

    def save_result(self, document_id: str, result: DocumentResult) -> None:
        row = self.get(document_id)
        if row is None:
            raise KeyError(document_id)
        if row["state"] == "approved":
            raise ValueError("Approved documents cannot be processed again")
        state = "failed" if result.error else "review"
        with self.db:
            self.db.execute(
                "UPDATE documents SET kind=?,state=?,result_json=?,edited_json=?,updated_at=? WHERE id=?",
                (result.document_type, state, result.model_dump_json(),
                 json.dumps(result.extracted, ensure_ascii=False) if result.extracted else None,
                 _now(), document_id),
            )
            self._event(document_id, "processed", {"status": result.status.value, "kind": result.document_type})

    def save_edits(self, document_id: str, data: dict) -> None:
        row = self.get(document_id)
        if row is None or not row["result_json"]:
            raise ValueError("Process this document first")
        if row["state"] == "approved":
            raise ValueError("Approved documents cannot be edited")
        if not isinstance(data, dict):
            raise TypeError("Edited data must be an object")
        before = json.loads(row["edited_json"] or "{}")
        changed = _changes(before, data)
        with self.db:
            self.db.execute("UPDATE documents SET edited_json=?,state='review',updated_at=? WHERE id=?",
                            (json.dumps(data, ensure_ascii=False), _now(), document_id))
            if changed:
                self._event(document_id, "edited", {"changes": changed, "actor": getpass.getuser()})

    def approve(self, document_id: str) -> list[str]:
        row = self.get(document_id)
        if row is None or not row["result_json"] or not row["edited_json"]:
            raise ValueError("Process this document first")
        result = DocumentResult.model_validate_json(row["result_json"])
        if result.document_type not in ("invoice", "receipt"):
            raise ValueError("Only invoices and receipts can be approved here")
        spec = catalog.get_schema(result.document_type)
        assert spec is not None
        try:
            document = spec.model.model_validate(json.loads(row["edited_json"]))
        except ValidationError as exc:
            return [f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors()]
        issues = validate(document, spec=spec)
        errors = [f"{issue.field}: {issue.message}" for issue in issues if issue.severity == "error"]
        if errors:
            return errors
        with self.db:
            self.db.execute("UPDATE documents SET state='approved',approved_at=?,updated_at=? WHERE id=?",
                            (_now(), _now(), document_id))
            self._event(document_id, "approved", {"actor": getpass.getuser()})
        return []

    def delete(self, document_id: str) -> None:
        row = self.get(document_id)
        if row is None:
            return
        with self.db:
            self.db.execute("DELETE FROM documents WHERE id=?", (document_id,))
        Path(row["stored_path"]).unlink(missing_ok=True)

    def backup(self, destination: Path) -> None:
        destination = Path(destination)
        if destination.resolve().is_relative_to(self.root.resolve()):
            raise ValueError("Save the backup outside the app data folder")
        from tempfile import TemporaryDirectory

        with TemporaryDirectory() as temp:
            snapshot = Path(temp) / "documents.sqlite3"
            with sqlite3.connect(snapshot) as target:
                self.db.backup(target)
            with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.write(snapshot, "documents.sqlite3")
                for file in self.files.iterdir():
                    if file.is_file():
                        archive.write(file, f"files/{file.name}")
        if os.name != "nt":
            destination.chmod(0o600)

    def restore_backup(self, archive_path: Path) -> None:
        """Restore a verified archive, retaining the previous store beside it."""
        with tempfile.TemporaryDirectory(dir=self.root.parent) as temp:
            staged = Path(temp) / "restored"
            staged.mkdir(mode=0o700)
            files = staged / "files"
            files.mkdir(mode=0o700)
            with zipfile.ZipFile(archive_path) as archive:
                entries = archive.infolist()
                names = {entry.filename for entry in entries}
                if len(entries) != len(names) or sum(entry.file_size for entry in entries) > 4_000_000_000:
                    raise ValueError("Backup has duplicate names or is too large")
                if "documents.sqlite3" not in names:
                    raise ValueError("Backup has no document database")
                for name in names:
                    if name == "documents.sqlite3" or name == "files/":
                        continue
                    if not name.startswith("files/") or not re.fullmatch(r"[0-9a-f]{64}\.[a-z0-9]+", name[6:]):
                        raise ValueError("Backup contains an unexpected file")
                for name in names:
                    if name.endswith("/"):
                        continue
                    target = staged / name
                    with archive.open(name) as source, target.open("wb") as output:
                        shutil.copyfileobj(source, output)
                    if os.name != "nt":
                        target.chmod(0o600)
            with sqlite3.connect(staged / "documents.sqlite3") as restored:
                if restored.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                    raise ValueError("Backup database failed integrity check")
                try:
                    rows = restored.execute("SELECT id,stored_path FROM documents").fetchall()
                except sqlite3.DatabaseError as exc:
                    raise ValueError("Backup is not a Docket Desktop database") from exc
                for document_id, stored_path in rows:
                    filename = Path(stored_path).name
                    if not filename.startswith(document_id + ".") or not (files / filename).is_file():
                        raise ValueError("Backup is missing a source document")
                    checksum = hashlib.sha256()
                    with (files / filename).open("rb") as source:
                        while block := source.read(1 << 20):
                            checksum.update(block)
                    digest = checksum.hexdigest()
                    if digest != document_id:
                        raise ValueError("Backup source document failed its checksum")
                    restored.execute("UPDATE documents SET stored_path=? WHERE id=?",
                                     (str(self.files / filename), document_id))
                restored.commit()

            previous = self.root.with_name(self.root.name + "-before-restore-" + datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S"))
            if previous.exists():
                raise FileExistsError(previous)
            self.close()
            try:
                self.root.rename(previous)
                staged.rename(self.root)
            except OSError:
                if not self.root.exists() and previous.exists():
                    previous.rename(self.root)
                DesktopStore.__init__(self, self.root)
                raise
            DesktopStore.__init__(self, self.root)

    def _approved_results(self) -> list[tuple[dict[str, Any], DocumentResult]]:
        approved = [row for row in self.list_documents() if row["state"] == "approved"]
        results: list[tuple[dict[str, Any], DocumentResult]] = []
        for row in approved:
            result = DocumentResult.model_validate_json(row["result_json"])
            data = json.loads(row["edited_json"])
            results.append((row, result.model_copy(update={"extracted": data})))
        return results

    def export_approved(self, summary: Path) -> tuple[int, Path]:
        results = self._approved_results()
        items_path = summary.with_name(summary.stem + "-items.csv")
        fields_path = summary.with_name(summary.stem + "-fields.csv")
        processing_path = summary.with_name(summary.stem + "-processing.csv")
        history_path = summary.with_name(summary.stem + "-history.csv")
        with Path(summary).open("w", encoding="utf-8-sig", newline="") as out:
            writer = csv.DictWriter(out, DOCUMENT_COLUMNS)
            writer.writeheader()
            for document, result in results:
                values = _document_row(document, result)
                writer.writerow({k: _csv_safe(v) for k, v in values.items()})
        with items_path.open("w", encoding="utf-8-sig", newline="") as out:
            writer = csv.DictWriter(out, ITEM_COLUMNS)
            writer.writeheader()
            for _, result in results:
                for row in line_item_rows(result):
                    writer.writerow({k: _csv_safe(v) for k, v in row.items()})
        for destination, records in (
            (fields_path, ((row["id"], result.extracted) for row, result in results)),
            (processing_path, ((row["id"], json.loads(row["result_json"])) for row, _ in results)),
        ):
            with destination.open("w", encoding="utf-8-sig", newline="") as out:
                writer = csv.DictWriter(out, DETAIL_COLUMNS)
                writer.writeheader()
                for document_id, data in records:
                    for item in _detail_rows(document_id, data):
                        writer.writerow({key: _csv_safe(value) for key, value in item.items()})
        with history_path.open("w", encoding="utf-8-sig", newline="") as out:
            writer = csv.DictWriter(out, HISTORY_COLUMNS)
            writer.writeheader()
            for document, _ in results:
                for event in self.history(document["id"]):
                    detail = event["detail_json"] or "{}"
                    for index, start in enumerate(range(0, max(1, len(detail)), MAX_CELL_CHARS), start=1):
                        values = {"document_id": document["id"], "revision_id": event["id"],
                                  "at": event["at"], "action": event["action"], "part": index,
                                  "detail_json": detail[start:start + MAX_CELL_CHARS]}
                        writer.writerow({key: _csv_safe(value) for key, value in values.items()})
        if os.name != "nt":
            for path in (Path(summary), items_path, fields_path, processing_path, history_path):
                path.chmod(0o600)
        return len(results), items_path

    def export_approved_xlsx(self, destination: Path) -> int:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        results = self._approved_results()
        workbook = Workbook()
        documents = workbook.active
        assert documents is not None
        documents.title = "Documents"
        items = workbook.create_sheet("Line items")
        fields = workbook.create_sheet("Fields")
        processing = workbook.create_sheet("Processing")
        history = workbook.create_sheet("History")

        def fill_sheet(sheet, columns: tuple[str, ...], rows) -> None:
            sheet.append(columns)
            for values in rows:
                sheet.append([_xlsx_safe(values.get(column)) for column in columns])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0E786E")
            for index, column in enumerate(columns, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = min(42, max(14, len(column) + 2))

        fill_sheet(documents, DOCUMENT_COLUMNS,
                   (_document_row(row, result) for row, result in results))
        fill_sheet(items, ITEM_COLUMNS,
                   (item for _, result in results for item in line_item_rows(result)))
        fill_sheet(fields, DETAIL_COLUMNS,
                   (item for row, result in results for item in _detail_rows(row["id"], result.extracted)))
        fill_sheet(processing, DETAIL_COLUMNS,
                   (item for row, _ in results for item in _detail_rows(
                       row["id"], json.loads(row["result_json"]))))
        fill_sheet(history, HISTORY_COLUMNS,
                   ({"document_id": row["id"], "revision_id": event["id"],
                     "at": event["at"], "action": event["action"], "part": index,
                     "detail_json": detail[start:start + MAX_CELL_CHARS]}
                    for row, _ in results for event in self.history(row["id"])
                    for detail in [event["detail_json"] or "{}"]
                    for index, start in enumerate(range(0, max(1, len(detail)), MAX_CELL_CHARS), start=1)))
        workbook.save(destination)
        if os.name != "nt":
            Path(destination).chmod(0o600)
        return len(results)

    def export_approved_json(self, destination: Path) -> int:
        results = self._approved_results()
        payload = {
            "format_version": 2,
            "documents": [
                {
                    "document_id": row["id"],
                    "filename": row["filename"],
                    "stored_path": row["stored_path"],
                    "document_type": result.document_type,
                    "schema_id": result.schema_id,
                    "schema_version": result.schema_version,
                    "approval_status": "approved",
                    "created_at": row["created_at"],
                    "updated_at": row["updated_at"],
                    "approved_at": row["approved_at"],
                    "extracted": result.extracted,
                    "processing_result": json.loads(row["result_json"]),
                    "history": [
                        {"revision_id": event["id"], "at": event["at"],
                         "action": event["action"], "detail": json.loads(event["detail_json"] or "{}")}
                        for event in self.history(row["id"])
                    ],
                }
                for row, result in results
            ],
        }
        with Path(destination).open("w", encoding="utf-8") as out:
            json.dump(payload, out, ensure_ascii=False, indent=2)
            out.write("\n")
        if os.name != "nt":
            Path(destination).chmod(0o600)
        return len(results)
